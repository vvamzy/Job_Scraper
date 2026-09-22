#!/usr/bin/env python3
"""
Job Scraper for Naukri.com

Extracts job listings matching a desired experience range and saves them to an
Excel workbook (.xlsx) with clickable job-title hyperlinks.

Usage:
    python job_scraper.py                     # defaults: python jobs, pages 1-5, 4 yrs exp
    python job_scraper.py --query "data engineer" --pages 1 3 --experience 4
    python job_scraper.py --query "devops" --pages 1 10 --output devops_jobs.xlsx --headless
"""

from __future__ import annotations

import argparse
import random
import re
import sys
import time
from dataclasses import dataclass, field

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait
from webdriver_manager.chrome import ChromeDriverManager

# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------

BASE_URL = "https://www.naukri.com"


@dataclass
class Job:
    title: str = ""
    company: str = ""
    rating: str = ""
    experience: str = ""
    location: str = ""
    requirements: str = ""
    tech_stack: str = ""
    link: str = ""

    def is_complete(self) -> bool:
        """A listing is usable when we at least know the title and where it links to."""
        return bool(self.title and self.link)


# ---------------------------------------------------------------------------
# Selenium helpers
# ---------------------------------------------------------------------------


def build_driver(headless: bool = False) -> webdriver.Chrome:
    """Create a Chrome driver configured to look like a normal human browser."""
    options = Options()
    options.add_argument("--start-maximized")
    options.add_argument("--disable-blink-features=AutomationControlled")
    options.add_argument("--no-first-run")
    options.add_argument("--no-default-browser-check")
    options.add_argument("--disable-popup-blocking")
    # A normal, current Chrome user agent (keeps Naukri's bot detector happy-ish)
    options.add_argument(
        "--user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36"
    )
    if headless:
        options.add_argument("--headless=new")
        options.add_argument("--window-size=1920,1080")
    # Hide the "Chrome is being controlled by automated software" bar
    options.add_experimental_option("excludeSwitches", ["enable-automation"])
    options.add_experimental_option("useAutomationExtension", False)
    options.add_experimental_option("prefs", {
        "credentials_enable_service": False,
        "profile.password_manager_enabled": False,
    })

    service = Service(ChromeDriverManager().install())
    driver = webdriver.Chrome(service=service, options=options)

    # Neutralise the navigator.webdriver flag so Selenium isn't detected
    driver.execute_cdp_cmd("Page.addScriptToEvaluateOnNewDocument", {
        "source": (
            "Object.defineProperty(navigator, 'webdriver', {get: () => undefined});"
            "window.chrome = window.chrome || {runtime: {}};"
        )
    })
    return driver


def human_pause(driver: webdriver.Chrome) -> None:
    """Scroll a bit and pause like a person, which reduces bot detection."""
    driver.execute_script("window.scrollBy(0, 600);")
    time.sleep(random.uniform(1.2, 2.4))


def dismiss_popups(driver: webdriver.Chrome) -> None:
    """Click through the cookie / 'interested?' popups Naukri likes to show."""
    candidates = [
        (By.ID, "wiz-close-notification-icon"),
        (By.ID, "close"),
        (By.XPATH, '//*[contains(text(), "Got it")]'),
        (By.XPATH, '//*[contains(text(), "GOT IT")]'),
        (By.XPATH, '//button[contains(., "I am Interested")]'),
        (By.XPATH, '//button[contains(., "Accept")]'),
        (By.XPATH, '//*[contains(@id, "dismiss")]'),
    ]
    for by, selector in candidates:
        try:
            el = driver.find_element(by, selector)
            if el.is_displayed():
                el.click()
                time.sleep(0.6)
        except Exception:
            continue


def page_url(query: str, page: int) -> str:
    """Build a Naukri search URL for the given query and page number."""
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", query.strip().lower()).strip("-")
    url = f"{BASE_URL}/{slug}-jobs"
    if page > 1:
        url += f"?pageNo={page}"
    return url


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------


def _text(el) -> str:
    return (el.get_text(" ", strip=True) if el else "")


def _rating_of(card) -> str:
    """Rating like '4.2' from a card, or '' if missing."""
    for sel in ("a.rating .main-2", ".rating .main-2", "span.rating", ".rating"):
        m = re.search(r"(\d+(?:\.\d+)?)", _text(card.select_one(sel)))
        if m:
            return m.group(1)
    # Fallback: a decimal rating (e.g. '4.5') anywhere in the card;
    # requiring the decimal avoids matching stray numbers like '1 day ago'.
    m = re.search(r"(?<![\d.])([1-5]\.\d+)(?![\d.])", card.get_text(" "))
    return m.group(1) if m else ""


def _experience_of(card) -> str:
    """Experience like '6-11 Yrs' or 'Fresher' from a card."""
    for sel in (".exp-wrap .expwdth", "span.exp", ".experience",
                "span.icl-Exp", "[class*='experience']"):
        val = _text(card.select_one(sel))
        if val:
            return val
    return ""


def _location_of(card) -> str:
    for sel in (".loc-wrap .locWdth", "span.loc", ".location",
                "span.icl-Location", "[class*='location']"):
        val = _text(card.select_one(sel))
        if val:
            return val
    return ""


def _skills_of(card) -> str:
    """Tech stack / skills joined into a comma-separated string."""
    skills = []
    for el in card.select(".tag-li, .skill, .tag, .tags, [class*='skill']"):
        t = _text(el)
        if t:
            skills.append(t)
    return ", ".join(dict.fromkeys(skills))


def _requirements_of(card) -> str:
    """Short description / minimum requirements from the card."""
    for sel in (".job-description", ".desc", "li[class*='desc']", ".job-desc"):
        val = _text(card.select_one(sel))
        if val:
            return val
    return ""


def parse_card(card) -> Job:
    job = Job()

    # --- Title + link ------------------------------------------------------
    anchors = [a for a in card.find_all("a", href=True) if a["href"].startswith("http")]
    preferred = [a for a in anchors if "job-listings" in a["href"]]
    for a in preferred or anchors:
        title = _text(a)
        if title and len(title) > 1:
            job.title = title
            job.link = a["href"]
            break

    # --- Company ------------------------------------------------------------
    for sel in ("a.comp-name", ".comp-name", "span.companyName", ".company-name",
                "a.subTitle", ".subTitle", "a.company"):
        val = _text(card.select_one(sel))
        if val:
            job.company = val
            break
    if not job.company:
        joined = card.get_text(" ")
        m = re.search(r"at\s+([A-Za-z0-9][A-Za-z0-9 &.\-]{1,60})", joined)
        if m:
            job.company = m.group(1).strip()

    job.rating = _rating_of(card)
    job.experience = _experience_of(card)
    job.location = _location_of(card)
    job.tech_stack = _skills_of(card)
    job.requirements = _requirements_of(card)

    # --- Fallback: location often sits right after the company name ---------
    if not job.location and job.company:
        joined = card.get_text(" ")
        m = re.search(
            re.escape(job.company) + r"\s*([A-Z][A-Za-z ,\-]*?)(?:\d+(?:\.\d+)?)?\s*(?:Yrs?|Years?)",
            joined,
        )
        if m:
            job.location = m.group(1).strip().rstrip(",")
    return job


def experience_fits(exp_text: str, target: int) -> bool:
    """
    Does `target` years of experience fit within the listed range?

    Examples (target = 4):
      "3-5 Yrs"  -> True     "4-8 Yrs" -> True
      "5-10 Yrs" -> False    "2-6 Yrs" -> True
      "Fresher"  -> True only when target == 0
      "4 Yrs"    -> True when target == 4
    """
    text = exp_text.lower()
    if not text:
        return True  # unknown range -> keep it, let the user decide
    if "fresher" in text or "entry" in text:
        return target == 0
    nums = [int(n) for n in re.findall(r"\d+", text)]
    if not nums:
        return True
    low, high = nums[0], (nums[1] if len(nums) > 1 else nums[0])
    return low <= target <= high


# ---------------------------------------------------------------------------
# Scraping
# ---------------------------------------------------------------------------


def scrape_pages(driver, query: str, start_page: int, end_page: int) -> list[Job]:
    jobs: dict[str, Job] = {}
    wait = WebDriverWait(driver, 20)

    for page in range(start_page, end_page + 1):
        url = page_url(query, page)
        print(f"[*] Fetching page {page}: {url}")
        driver.get(url)
        human_pause(driver)
        dismiss_popups(driver)

        found = 0
        try:
            # Wait for job cards to render (classic or current Naukri markup)
            wait.until(
                EC.presence_of_element_located(
                    (By.CSS_SELECTOR,
                     ".srp-jobtuple-wrapper, div.jobTuple, "
                     "[class*='job-card'], [class*='jobCard']")
                )
            )
            time.sleep(2)
        except Exception:
            print(f"[!] No job cards found on page {page}. "
                  f"Dumping HTML to page_{page}_debug.html for inspection.")
            with open(f"page_{page}_debug.html", "w", encoding="utf-8") as fh:
                fh.write(driver.page_source)
            continue

        soup = BeautifulSoup(driver.page_source, "html.parser")
        cards = (
            soup.select(".srp-jobtuple-wrapper")
            or soup.select("div.jobTuple")
            or soup.select("[class*='job-card']")
            or soup.select("[class*='jobCard']")
        )
        for card in cards:
            job = parse_card(card)
            if not job.is_complete():
                continue
            if not experience_fits(job.experience, args.experience):
                continue
            if job.link in jobs:
                continue
            jobs[job.link] = job
            found += 1

        print(f"    -> collected {found} matching job(s) on this page "
              f"({len(jobs)} total)")

    return list(jobs.values())


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------


def export_excel(jobs: list[Job], output: str) -> None:
    headers = [
        "Job Title",
        "Company Name",
        "Rating",
        "Experience",
        "Location",
        "Minimum Requirements",
        "Tech Stack",
        "Job Link",
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row, job in enumerate(jobs, start=2):
        title_cell = ws.cell(row=row, column=1, value=job.title)
        title_cell.hyperlink = job.link            # <-- makes the title clickable
        title_cell.font = Font(color="0563C1", underline="single")
        ws.cell(row=row, column=2, value=job.company)
        ws.cell(row=row, column=3, value=job.rating)
        ws.cell(row=row, column=4, value=job.experience)
        ws.cell(row=row, column=5, value=job.location)
        ws.cell(row=row, column=6, value=job.requirements)
        ws.cell(row=row, column=7, value=job.tech_stack)
        ws.cell(row=row, column=8, value=job.link).hyperlink = job.link
        for col in range(1, len(headers) + 1):
            ws.cell(row=row, column=col).border = border

    widths = [45, 25, 8, 12, 30, 60, 35, 60]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(output)
    print(f"[+] Saved {len(jobs)} job(s) to {output}")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> None:
    parser = argparse.ArgumentParser(description="Scrape job listings from Naukri.com")
    parser.add_argument("--query", default="python",
                        help="Job keyword/search term (default: python)")
    parser.add_argument("--pages", nargs=2, type=int, default=[1, 5], metavar=("START", "END"),
                        help="Page range to scrape (default: 1 5)")
    parser.add_argument("--experience", type=int, default=4,
                        help="Your years of experience; listings must cover this "
                             "range (default: 4)")
    parser.add_argument("--output", default="filtered_jobs.xlsx",
                        help="Output Excel file (default: filtered_jobs.xlsx)")
    parser.add_argument("--headless", action="store_true",
                        help="Run Chrome without a visible window")
    global args
    args = parser.parse_args()

    start_page, end_page = args.pages
    if end_page < start_page:
        print("[!] End page must be >= start page.")
        sys.exit(1)

    driver = None
    try:
        print(f"[*] Starting Chrome for '{args.query}' (pages {start_page}-{end_page}, "
              f"experience fit: {args.experience} yrs)")
        driver = build_driver(headless=args.headless)
        jobs = scrape_pages(driver, args.query, start_page, end_page)
        if jobs:
            export_excel(jobs, args.output)
        else:
            print("[!] No matching jobs found. Try a different query or widen "
                  "the page range.")
    finally:
        if driver:
            driver.quit()


if __name__ == "__main__":
    main()